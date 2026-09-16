
from __future__ import annotations

import re
import uuid
import logging
from pathlib import Path
from typing import List, Dict, Any, Tuple

try:
    from agents.rag_agent.doc_parser import MedicalDocParser as _MedicalDocParser
except ImportError:
    _MedicalDocParser = None

from core.errors import MIHCError

logger = logging.getLogger(__name__)

class IngestionPipeline:

    def __init__(self, config, embedder, milvus_store, keyword_store, database):
        self.config = config
        self.embedder = embedder
        self.milvus = milvus_store
        self.keyword = keyword_store
        self.db = database
        self.parser = _MedicalDocParser() if _MedicalDocParser else None

    def parse(self, file_path: str) -> Tuple[str, str]:
        path = Path(file_path)
        if not path.exists():
            raise MIHCError(f"文件不存在: {file_path}", code="file_not_found", status_code=404)
        if path.suffix.lower() not in (".pdf", ".docx", ".md", ".txt"):
            raise MIHCError(f"暂不支持的文件类型: {path.suffix}", code="unsupported_file", status_code=422)
        logger.info("Parsing document: %s", file_path)
        try:
            doc, _images = self.parser.parse_document(str(path), "./data/parsed_docs")
            markdown = doc.export_to_markdown()
        except Exception as exc:
            logger.warning("Docling 解析失败（%s），回退轻量解析器", exc)
            markdown = self._fallback_parse(path)
        return markdown, path.name

    @staticmethod
    def _fallback_parse(path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"## 第 {i} 页\n{text}")
            return "\n\n".join(pages)
        if suffix == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(str(path))
            parts = []
            for para in doc.paragraphs:
                if para.style.name.startswith("Heading"):
                    parts.append(f"## {para.text}")
                elif para.text.strip():
                    parts.append(para.text)
            for table in doc.tables:
                parts.append("\n".join(" | ".join(c.text for c in row.cells) for row in table.rows))
            return "\n\n".join(parts)
        return path.read_text(encoding="utf-8", errors="ignore")

    @staticmethod
    def clean(text: str) -> str:
        text = re.sub(r"\n\s*\d+\s*\n", "\n", text)
        text = re.sub(r"[ \t]{2,}", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @staticmethod
    def chunk_by_sections(text: str, max_chars: int = 512, overlap_chars: int = 50) -> List[Dict[str, str]]:
        sections = re.split(r"(?=\n#+\s)", text)
        chunks: List[Dict[str, str]] = []
        current = ""
        current_title = ""

        def flush():
            nonlocal current, current_title
            if current.strip():
                chunks.append({"text": current.strip(), "title": current_title})
            current = ""
            current_title = ""

        for sec in sections:
            sec = sec.strip()
            if not sec:
                continue
            title_match = re.match(r"#+\s+(.*)", sec.split("\n", 1)[0])
            title = title_match.group(1).strip() if title_match else ""
            paras = re.split(r"\n(?=\S)", sec)
            for para in paras:
                para = para.strip()
                if not para:
                    continue
                if current and len(current) + len(para) > max_chars:
                    flush()
                if not current:
                    current_title = title
                current = (current + "\n\n" + para).strip()
        flush()
        if overlap_chars > 0 and len(chunks) > 1:
            for i in range(1, len(chunks)):
                chunks[i]["text"] = chunks[i - 1]["text"][-overlap_chars:] + "\n" + chunks[i]["text"]
        return chunks

    def ingest(self, file_path: str, *, version: str = "", source: str = "",
               permission: str = "research_team", tenant_id: str = "mihc") -> Dict[str, Any]:
        markdown, file_name = self.parse(file_path)
        text = self.clean(markdown)
        sections = self.chunk_by_sections(text)

        doc_id = uuid.uuid4().hex
        source = source or str(file_path)
        embed_model = self.config.embedding.model_name

        chunk_records: List[Dict[str, Any]] = []
        for i, sec in enumerate(sections):
            chunk_id = f"{doc_id}-chunk-{i:03d}"
            chunk_records.append({
                "chunk_id": chunk_id,
                "doc_id": doc_id,
                "text": sec["text"],
                "title": sec.get("title", ""),
                "source": source,
                "version": version,
                "tenant_id": tenant_id,
                "embedding_model": embed_model,
            })

        embeddings = self.embedder.embed_documents([c["text"] for c in chunk_records])

        try:
            self.milvus.upsert_chunks(chunk_records, embeddings)
        except Exception as exc:
            logger.warning("Milvus upsert degraded (%s), keyword index only", exc)
        self.keyword.upsert_chunks(chunk_records)

        self.db.add_document(doc_id=doc_id, file_name=file_name, source=source, version=version,
                             tenant_id=tenant_id, permission=permission,
                             embedding_model=embed_model, chunk_count=len(chunk_records))
        self.db.add_chunks([{"chunk_id": c["chunk_id"], "doc_id": c["doc_id"],
                             "title": c["title"], "text": c["text"],
                             "source": c["source"], "version": c["version"],
                             "tenant_id": c["tenant_id"], "embedding_model": c["embedding_model"]}
                            for c in chunk_records])

        logger.info("Ingested %s: %d chunks -> Milvus + keyword store + PG", file_name, len(chunk_records))
        return {
            "doc_id": doc_id,
            "file_name": file_name,
            "chunks": len(chunk_records),
            "embedding_model": embed_model,
        }
